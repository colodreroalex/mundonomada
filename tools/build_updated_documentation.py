from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / 'docs' / 'Mundo_Nomada_Documentacion_Actualizada.docx'


def set_cell_shading(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), fill)
    props.append(shading)


def set_cell_border(cell):
    props = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        tag = OxmlElement(f'w:{edge}')
        tag.set(qn('w:val'), 'single')
        tag.set(qn('w:sz'), '4')
        tag.set(qn('w:color'), 'D9D9D9')
        borders.append(tag)
    props.append(borders)


def add_heading(doc, text, level=1):
    paragraph = doc.add_heading(text, level=level)
    for run in paragraph.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
    paragraph.paragraph_format.space_before = Pt(14 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(6)
    return paragraph


def add_body(doc, text):
    paragraph = doc.add_paragraph(text)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.15
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(text, style='List Bullet')
    paragraph.paragraph_format.space_after = Pt(3)
    return paragraph


def add_table(doc, rows):
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    header = table.rows[0].cells
    for index, value in enumerate(('Entidad', 'Información principal', 'Reglas de integridad')):
        header[index].text = value
        set_cell_shading(header[index], '1F4E78')
        for run in header[index].paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.bold = True
        header[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_border(header[index])
    for index, values in enumerate(rows):
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    run.font.size = Pt(9)
        if index % 2 == 1:
            for cell in cells:
                set_cell_shading(cell, 'EAF2F8')
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal = doc.styles['Normal']
    normal.font.name = 'Aptos'
    normal._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
    normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    normal.font.size = Pt(10.5)

    title = doc.add_paragraph(style='Title')
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run('Mundo Nomada Documentacion Tecnica Actualizada')
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.font.name = 'Aptos Display'
    run.font.size = Pt(24)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run('Estado del proyecto y base de datos para puesta en produccion')
    subtitle_run.font.size = Pt(12)
    subtitle_run.font.color.rgb = RGBColor(0, 0, 0)
    subtitle.paragraph_format.space_after = Pt(20)

    add_body(doc, 'Este documento sustituye las afirmaciones de la memoria original que no se corresponden con el codigo revisado. Sirve como guia de desarrollo y de puesta en marcha: distingue lo que ya existe de lo que debe completarse antes de aceptar pedidos reales.')

    add_heading(doc, 'Alcance y estado actual')
    add_body(doc, 'Mundo Nomada es una tienda online para catalogo, clientes, carrito y pedidos. El frontend esta desarrollado en Angular y el backend es una API PHP. El proyecto se mantiene en GitHub con ramas y Pull Requests para que cada cambio pueda revisarse antes de llegar a main.')
    add_bullet(doc, 'Implementado en el codigo: catalogo, categorias, registro e inicio de sesion, perfil, administracion, carrito y estructura de pedidos.')
    add_bullet(doc, 'En revision: el Pull Request numero 1 protege la propiedad del carrito y endurece el checkout. Debe fusionarse y probarse antes del despliegue.')
    add_bullet(doc, 'Pendiente: pago de produccion, pedidos de invitado, correo transaccional, pruebas automatizadas y despliegue publico.')

    add_heading(doc, 'Arquitectura')
    add_body(doc, 'El navegador Angular se comunica unicamente con la API PHP. La API valida la sesion y las entradas antes de usar la base de datos. En produccion no se exponen credenciales de base de datos, ni se conectan los navegadores directamente a las tablas de negocio.')
    add_bullet(doc, 'Desarrollo actual: Angular, PHP y MySQL MariaDB local mediante XAMPP.')
    add_bullet(doc, 'Destino de produccion: el proyecto Supabase PostgreSQL mundo-nomada ya existe en West EU Ireland. La migracion inicial versionada se aplico y las tablas de negocio tienen RLS activado.')
    add_bullet(doc, 'Archivos privados: la conexion se define por variables de entorno o por archivos locales que Git ignora. Para Supabase, PHP necesita la extension PDO PostgreSQL activada.')

    add_heading(doc, 'Modelo de datos')
    add_body(doc, 'La fuente de verdad del modelo es database schema sql para el entorno MySQL local y supabase migrations para PostgreSQL. Ambas definiciones no incluyen datos de clientes ni productos reales.')
    add_table(doc, [
        ('users', 'Cuenta, correo, hash de contrasena, rol y sesiones persistentes.', 'Correo unico. Contrasenas hasheadas. Roles admin y user.'),
        ('categorias', 'Nombre y descripcion de cada categoria.', 'Nombre unico.'),
        ('productos', 'Nombre, precio, stock, categoria, color, talla e imagen.', 'Precio y stock no negativos. Categoria valida.'),
        ('carrito', 'Producto y cantidad que un usuario desea comprar.', 'Una linea por usuario y producto. Cantidad positiva.'),
        ('orders', 'Pedido, total, estado, pago y datos de envio.', 'Importes decimales. Referencia de pago unica si existe.'),
        ('order_items', 'Detalle y precio historico de cada linea del pedido.', 'Cantidad positiva y relacion obligatoria con pedido y producto.'),
    ])

    add_heading(doc, 'Reglas de seguridad de la base de datos')
    add_bullet(doc, 'Los importes usan tipos decimales, no float, para evitar errores de redondeo en dinero.')
    add_bullet(doc, 'Las claves foraneas impiden carritos y pedidos sin usuario o producto valido.')
    add_bullet(doc, 'El carrito tiene una restriccion unica de usuario y producto; evita duplicados y facilita actualizar cantidades.')
    add_bullet(doc, 'Los datos de produccion no se suben al repositorio. Las copias de phpMyAdmin y los volcados con clientes se excluyen de Git.')
    add_bullet(doc, 'Supabase tendra Row Level Security activado sin politicas de acceso directo. El backend sera la unica capa autorizada hasta implementar una autenticacion especifica.')

    add_heading(doc, 'Configuracion local')
    add_body(doc, 'Para ejecutar una copia local, se crea una base vacia con database schema sql y se configura el archivo database local php a partir del ejemplo incluido en el backend. La cuenta de la aplicacion debe ser distinta de root y tener solo permisos sobre la base de Mundo Nomada.')
    add_body(doc, 'Para Supabase, la migracion inicial ya se aplico desde el archivo versionado. Antes de usar la CLI para cambios futuros, se debe registrar esa migracion como aplicada, y despues cada cambio se aplicara desde un nuevo archivo versionado. No se editaran tablas manualmente en el panel remoto porque se perderia la trazabilidad de los cambios.')

    add_heading(doc, 'Guia manual de conexion a Supabase')
    add_bullet(doc, 'En Project Settings y Database, restablece la contrasena de PostgreSQL y guardala en un gestor de contrasenas. No la compartas por chat ni la subas a Git.')
    add_bullet(doc, 'En Connect, usa Session pooler. Copia sus datos en config supabase local php dentro del backend a partir del ejemplo y completa la contrasena. Ese archivo esta ignorado por Git.')
    add_bullet(doc, 'En C XAMPP php php ini, habilita pdo pgsql y pgsql, y reinicia Apache. Sin esos drivers PHP no puede conectarse a PostgreSQL.')
    add_bullet(doc, 'Crea un token personal de Supabase en tu cuenta y ejecuta la reparacion del historial indicada en supabase README antes de usar db push. No compartas el token.')

    add_heading(doc, 'Pasos antes de vender')
    numbered = [
        'Fusionar el Pull Request de seguridad y probar que un usuario no puede consultar ni modificar el carrito de otro.',
        'Configurar la contrasena y los drivers PDO PostgreSQL, y adaptar la conexion PHP de mysqli a PDO PostgreSQL.',
        'Probar registro, inicio de sesion, carrito, stock, pedido y permisos de administrador con datos ficticios.',
        'Completar la pasarela de pago para que el backend confirme el pago con el proveedor antes de marcar un pedido como pagado.',
        'Configurar HTTPS, cookies seguras, CORS de produccion, limites de inicio de sesion, copias de seguridad y prueba de restauracion.',
        'Publicar las paginas legales reales, politica de privacidad y datos de contacto de la tienda.',
    ]
    for item in numbered:
        doc.add_paragraph(item, style='List Number')

    add_heading(doc, 'Funcionalidades futuras')
    add_body(doc, 'Microservicios, WebSockets, recomendaciones, puntos, varios idiomas, varias divisas, MFA, CI CD y escalado automatico son posibles mejoras, pero no se consideran implementadas ni bloquean la primera version. La prioridad es terminar un flujo de compra seguro, verificable y mantenible.')

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == '__main__':
    build()
